"""ФИО одним полем, готовые ID, конфликты и реальное окно создания связи."""

from decimal import Decimal
import os
from pathlib import Path
import tempfile
import unittest
import uuid

from excel2pg.entities import WorkbookModel
from excelprep.model import Project
from test_excelprep import table, parse_export


def people():
    users = table(
        "users",
        [
            ("user_id", "BIGINT"),
            ("last_name", "TEXT"),
            ("first_name", "TEXT"),
            ("middle_name", "TEXT"),
        ],
        [
            dict(
                user_id=17,
                last_name="Иванов",
                first_name="Иван",
                middle_name="Сергеевич",
            ),
            dict(
                user_id=48,
                last_name="Иванов",
                first_name="Пётр",
                middle_name="Алексеевич",
            ),
        ],
        ("user_id",),
    )
    orders = table(
        "orders",
        [("order_id", "BIGINT"), ("client_name", "TEXT"), ("user_id", "BIGINT")],
        [
            dict(
                order_id=1, client_name="  ИВАНОВ\u00a0Иван   Сергеевич ", user_id=None
            ),
            dict(order_id=2, client_name="Иванов Пётр Алексеевич", user_id=None),
        ],
        ("order_id",),
    )
    return Project(WorkbookModel([users, orders]))


def arguments():
    return dict(
        child_name="orders",
        columns=("client_name",),
        parent_name="users",
        match_columns=("last_name", "first_name", "middle_name"),
        reference_column="user_id",
        combine=True,
    )


class LinkTests(unittest.TestCase):
    def test_uuid_case_and_numeric_precision_ids_do_not_create_false_conflicts(self):
        guid = uuid.uuid4()
        p = Project(
            WorkbookModel(
                [
                    table("people", [("id", "UUID")], [dict(id=guid)], ("id",)),
                    table(
                        "records",
                        [("id", "BIGINT"), ("owner", "TEXT")],
                        [dict(id=1, owner=str(guid).upper())],
                        ("id",),
                    ),
                ]
            )
        )
        p.change(
            p.connect_tables,
            "records",
            ("owner",),
            "people",
            ("id",),
            "owner",
            mode="id",
        )
        self.assertEqual(p.table("records").rows[0]["owner"], guid)
        self.assertEqual(p.validate(), [])
        q = people()
        q.table("users").primary_key.pg_type = "NUMERIC(12,2)"
        q.table("orders").rows[0]["user_id"] = "17.00"
        q.table("orders").rows[1]["user_id"] = "48"
        q.change(
            q.connect_tables,
            "orders",
            ("user_id",),
            "users",
            ("user_id",),
            "user_id",
            mode="id",
        )
        self.assertEqual(q.validate(), [])

    def test_import_suffix_keeps_real_user_id_as_key(self):
        from openpyxl import Workbook
        from excelprep.model import load_excel

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "Users_import.xlsx"
            book = Workbook()
            book.active.append(["user_id", "Фамилия", "Имя"])
            book.active.append([17, "Иванов", "Иван"])
            book.save(path)
            model = load_excel([path])
        self.assertEqual(model.tables[0].name, "users")
        self.assertEqual(model.tables[0].primary_key.name, "user_id")
        self.assertEqual(model.tables[0].rows[0]["user_id"], 17)

    def test_fio_parts_normalization_existing_empty_id_and_preview_does_not_mutate(
        self,
    ):
        p = people()
        old = p.as_dict()
        report = p.preview_link(**arguments())
        self.assertEqual(report["matched"], 2)
        self.assertEqual(report["values"], [17, 48])
        self.assertEqual(p.as_dict(), old)
        p.change(p.connect_tables, **arguments())
        self.assertEqual([r["user_id"] for r in p.table("orders").rows], [17, 48])
        self.assertEqual(len(p.table("orders").columns), 3)
        self.assertEqual(p.validate(), [])
        p.undo()
        self.assertEqual(p.as_dict(), old)

    def test_homonyms_and_unknown_names_never_pick_random_user(self):
        p = people()
        p.table("users").rows.append(
            dict(
                user_id=77,
                last_name="Иванов",
                first_name="Иван",
                middle_name="Сергеевич",
            )
        )
        p.table("orders").rows[1]["client_name"] = "Неизвестный клиент"
        report = p.preview_link(**arguments())
        self.assertEqual((report["ambiguous"], report["missing"]), (1, 1))
        old = p.as_dict()
        with self.assertRaises(ValueError):
            p.change(p.connect_tables, **arguments())
        self.assertEqual(p.as_dict(), old)
        p.change(p.connect_tables, **arguments(), allow_missing=True)
        self.assertEqual([r["user_id"] for r in p.table("orders").rows], [None, None])
        self.assertEqual(p.validate(), [])

    def test_filled_ids_are_kept_and_conflicts_require_explicit_replacement(self):
        p = people()
        p.table("orders").rows[0].update(client_name="", user_id=17)
        p.table("orders").rows[1]["user_id"] = 17
        report = p.preview_link(**arguments())
        self.assertEqual(report["kept"], 1)
        self.assertEqual(report["conflicts"], 1)
        with self.assertRaises(ValueError):
            p.change(p.connect_tables, **arguments())
        p.change(p.connect_tables, **arguments(), replace_existing=True)
        # При явном полном пересчёте пустое ФИО не даёт произвольного пользователя.
        self.assertEqual([r["user_id"] for r in p.table("orders").rows], [None, 48])

    def test_id_types_numeric_and_text_codes_and_multiple_parents(self):
        p = people()
        p.table("orders").rows[0]["user_id"] = "17.00"
        p.table("orders").rows[1]["user_id"] = Decimal("48.0")
        p.change(
            p.connect_tables,
            "orders",
            ("user_id",),
            "users",
            ("user_id",),
            "user_id",
            mode="id",
        )
        self.assertEqual([r["user_id"] for r in p.table("orders").rows], [17, 48])
        stores = table(
            "stores",
            [("store_id", "TEXT")],
            [dict(store_id="001"), dict(store_id="002")],
            ("store_id",),
        )
        p.model.tables.append(stores)
        for c in p.table("orders").rows:
            c["store_code"] = "001"
        from excel2pg.entities import ColumnEntity

        p.table("orders").columns.append(
            ColumnEntity("store_code", "store_code", "TEXT")
        )
        p.change(
            p.connect_tables,
            "orders",
            ("store_code",),
            "stores",
            ("store_id",),
            "store_id",
            mode="id",
        )
        self.assertEqual(len(p.table("orders").foreign_keys), 2)
        self.assertEqual(p.validate(), [])
        p.table("orders").rows[0]["store_code"] = "1"
        self.assertEqual(
            p.preview_link(
                "orders",
                ("store_code",),
                "stores",
                ("store_id",),
                "new_store_id",
                mode="id",
            )["missing"],
            1,
        )

    def test_prepared_files_preserve_user_ids_and_relation(self):
        p = people()
        p.change(p.connect_tables, **arguments())
        with tempfile.TemporaryDirectory() as folder:
            model = parse_export(p.export(folder))
        orders = next(t for t in model.tables if t.name == "orders")
        self.assertEqual([r["user_id"] for r in orders.rows], [17, 48])
        self.assertTrue(orders.foreign_keys[0].enabled)


@unittest.skipUnless(os.getenv("DISPLAY"), "Нужен DISPLAY")
class LinkUiTests(unittest.TestCase):
    def test_choose_orders_users_fio_and_create_using_real_window(self):
        from excelprep.ui import ExcelPreparationApp

        app = ExcelPreparationApp()
        self.addCleanup(app.destroy)
        app.loaded(people())
        app.relation_dialog()
        app.update()
        window = app.links_window
        # Даже если сначала выбраны users, изменение первой таблицы на orders
        # автоматически выбирает users второй таблицей.
        window.child.set("orders")
        window.reset_fields()
        app.update()
        self.assertEqual(window.parent.get(), "users")
        self.assertEqual(window.mode.get(), "text")
        self.assertIn("ФИО целиком", window.parent_field.get())
        self.assertIn("Найдено: 2", window.summary.get())
        self.assertEqual(str(window.create_button.cget("state")), "normal")
        window.create_button.invoke()
        app.update()
        self.assertEqual(
            [r["user_id"] for r in app.project.table("orders").rows], [17, 48]
        )
        self.assertEqual(app.project.validate(), [])
        self.assertEqual(len(app.buttons), 4)


@unittest.skipUnless(os.getenv("CHUDO_TEST_DSN"), "Нужен CHUDO_TEST_DSN")
class LinkPostgresTests(unittest.TestCase):
    def test_orders_user_foreign_key_is_created_after_excel_export(self):
        import psycopg
        from excel2pg.postgres import import_model

        p = people()
        p.change(p.connect_tables, **arguments())
        schema = "linked_" + uuid.uuid4().hex
        dsn = os.environ["CHUDO_TEST_DSN"]
        try:
            with tempfile.TemporaryDirectory() as folder:
                result = import_model(parse_export(p.export(folder)), dsn, schema)
            self.assertEqual(result.foreign_keys_added, 1)
            with psycopg.connect(dsn, autocommit=True) as c:
                self.assertEqual(
                    c.execute(
                        f'SELECT user_id FROM "{schema}".orders ORDER BY order_id'
                    ).fetchall(),
                    [(17,), (48,)],
                )
                with self.assertRaises(psycopg.errors.ForeignKeyViolation):
                    c.execute(
                        f'UPDATE "{schema}".orders SET user_id=999 WHERE order_id=1'
                    )
        finally:
            with psycopg.connect(dsn, autocommit=True) as c:
                c.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
