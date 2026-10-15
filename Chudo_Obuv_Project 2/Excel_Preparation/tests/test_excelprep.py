"""Проверка разбиения, сохранности значений и импорта подготовленных файлов."""

from datetime import date, datetime, timezone
from decimal import Decimal
import json
import os
from pathlib import Path
import tempfile
import unittest
import uuid

from openpyxl import Workbook, load_workbook

from excel2pg.entities import ColumnEntity, TableEntity, WorkbookModel
from excelprep.model import (
    Project,
    english_name,
    load_excel,
    relations_sql,
    suggested_names,
)
from excel2pg.profile import read_import_profile


def table(name, columns, rows, key):
    return TableEntity(
        name,
        "test.xlsx",
        "Sheet1",
        [ColumnEntity(n, n, t, True, n in key) for n, t in columns],
        rows,
    )


def example():
    products = table(
        "products",
        [
            ("product_id", "TEXT"),
            ("product_name", "TEXT"),
            ("manufacturer", "TEXT"),
            ("country", "TEXT"),
            ("price", "NUMERIC"),
        ],
        [
            dict(
                product_id="001",
                product_name="Книга",
                manufacturer="Альфа",
                country="Россия",
                price=Decimal("350.50"),
            ),
            dict(
                product_id="002",
                product_name="Альбом",
                manufacturer="Альфа",
                country="Россия",
                price=Decimal("500.00"),
            ),
            dict(
                product_id="003",
                product_name="Тетрадь",
                manufacturer=None,
                country=None,
                price=Decimal("100.00"),
            ),
        ],
        ("product_id",),
    )
    return Project(WorkbookModel([products]))


def parse_export(path):
    from excel2pg.parser import parse_excel_files

    files = sorted(path.glob("*.xlsx"))
    profile = read_import_profile(path / "import_profile.json", files)
    return parse_excel_files(
        files,
        table_rules=profile["table_rules"],
        foreign_keys=profile["foreign_keys"],
        reference_rules=(),
        auto_foreign_keys=False,
    )


class PreparationTests(unittest.TestCase):
    def test_missing_keys_created_and_existing_ids_preserved(self):
        project = example()
        unkeyed = table(
            "categories",
            [("category_id", "TEXT"), ("name", "TEXT")],
            [dict(category_id="same", name="A"), dict(category_id="same", name="B")],
            (),
        )
        project.model.tables.append(unkeyed)
        created = project.prepare_keys()
        self.assertEqual(created, [("categories", "category_row_id")])
        self.assertEqual(project.table("categories").rows[0]["category_id"], "same")
        self.assertEqual([r["category_row_id"] for r in unkeyed.rows], [1, 2])
        self.assertEqual(project.table("products").primary_key.name, "product_id")
        self.assertEqual(project.prepare_keys(), [])
        self.assertEqual(project.validate(), [])
        self.assertEqual(
            suggested_names(project.table("products"), ("country",)),
            ("countries", "country_id"),
        )

    def test_translation_only_names_and_preserves_codes_null_words(self):
        self.assertEqual(english_name("Производители"), "manufacturers")
        self.assertEqual(english_name("Неизвестное поле"), "neizvestnoe_pole")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "Товары.xlsx"
            book = Workbook()
            book.active.append(["Код", "Наименование товара", "Комментарий"])
            book.active.append(["001", "NA", "=SUM(A1:A3)"])
            book.active["C2"].data_type = "s"
            book.active.append(["002", "NULL", None])
            book.save(path)
            model = load_excel([path])
            self.assertEqual(model.tables[0].name, "products")
            self.assertEqual(model.tables[0].rows[0]["id"], "001")
            self.assertEqual(model.tables[0].rows[0]["product_name"], "NA")
            self.assertEqual(model.tables[0].rows[1]["product_name"], "NULL")
            self.assertEqual(model.tables[0].rows[0]["kommentariy"], "=SUM(A1:A3)")

    def test_extract_repeated_entity_preserves_rows_and_nullable_link(self):
        project = example()
        project.change(
            project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
            ("manufacturer",),
        )
        parent = project.table("manufacturers")
        self.assertEqual(len(parent.rows), 1)
        self.assertEqual(len(project.table("products").rows), 3)
        self.assertEqual(
            [r["manufacturer_id"] for r in project.table("products").rows], [1, 1, None]
        )
        self.assertEqual(project.validate(), [])
        self.assertNotIn("manufacturer", project.table("products").rows[0])
        project.undo()
        self.assertEqual(project.model.table_names(), ["products"])
        self.assertIn("country", project.table("products").rows[0])

    def test_conflicting_dependency_rolls_back(self):
        project = example()
        project.table("products").rows[1]["country"] = "Китай"
        old = project.as_dict()
        with self.assertRaisesRegex(ValueError, "разные значения"):
            project.change(
                project.extract,
                "products",
                ("manufacturer", "country"),
                "manufacturers",
                "manufacturer_id",
                ("manufacturer",),
            )
        self.assertEqual(old, project.as_dict())

    def test_split_order_headers_and_move_user_fk(self):
        users = table(
            "users", [("user_id", "TEXT")], [dict(user_id="u1")], ("user_id",)
        )
        items = table(
            "order_items",
            [
                ("line_id", "BIGINT"),
                ("number", "BIGINT"),
                ("order_date", "DATE"),
                ("user_id", "TEXT"),
                ("product", "TEXT"),
            ],
            [
                dict(
                    line_id=1,
                    number=5,
                    order_date=date(2026, 9, 1),
                    user_id="u1",
                    product="A",
                ),
                dict(
                    line_id=2,
                    number=5,
                    order_date=date(2026, 9, 1),
                    user_id="u1",
                    product="B",
                ),
            ],
            ("line_id",),
        )
        project = Project(WorkbookModel([users, items]))
        project.add_relation("order_items", ("user_id",), "users", ("user_id",))
        project.change(
            project.extract,
            "order_items",
            ("number", "order_date", "user_id"),
            "orders",
            "order_id",
            ("number",),
        )
        self.assertEqual(len(project.table("orders").rows), 1)
        self.assertEqual(
            project.table("orders").foreign_keys[0].referenced_table, "users"
        )
        self.assertEqual(
            project.table("order_items").foreign_keys[0].referenced_table, "orders"
        )
        self.assertEqual(project.validate(), [])

    def test_composite_keys_orphans_and_rename_follow_links(self):
        parent = table(
            "editions",
            [("book", "TEXT"), ("edition", "BIGINT")],
            [dict(book="001", edition=1)],
            ("book", "edition"),
        )
        child = table(
            "sales",
            [("id", "BIGINT"), ("book", "TEXT"), ("edition", "BIGINT")],
            [dict(id=1, book="001", edition=1)],
            ("id",),
        )
        project = Project(WorkbookModel([parent, child]))
        project.add_relation(
            "sales", ("book", "edition"), "editions", ("book", "edition")
        )
        project.change(project.rename_column, "editions", "book", "isbn")
        project.change(project.rename_table, "editions", "versions")
        self.assertEqual(
            project.table("sales").foreign_keys[0].reference_columns,
            ("isbn", "edition"),
        )
        self.assertEqual(project.validate(), [])
        child.rows[0]["book"] = "absent"
        self.assertTrue(project.validate())

    def test_link_text_rejects_ambiguous_names_and_preserves_original(self):
        project = example()
        manufacturers = table(
            "manufacturers",
            [("manufacturer_id", "BIGINT"), ("name", "TEXT")],
            [
                dict(manufacturer_id=1, name="Альфа"),
                dict(manufacturer_id=2, name="Альфа"),
            ],
            ("manufacturer_id",),
        )
        project.model.tables.append(manufacturers)
        with self.assertRaisesRegex(ValueError, "Неоднозначное"):
            project.change(
                project.link_by_values,
                "products",
                ("manufacturer",),
                "manufacturers",
                ("name",),
                "manufacturer_id",
            )
        project.table("manufacturers").rows.pop()
        project.change(
            project.link_by_values,
            "products",
            ("manufacturer",),
            "manufacturers",
            ("name",),
            "manufacturer_id",
        )
        self.assertEqual(project.table("products").rows[0]["manufacturer"], "Альфа")
        self.assertEqual(project.table("products").rows[0]["manufacturer_id"], 1)
        self.assertEqual(project.validate(), [])

    def test_export_roundtrip_keys_types_and_literal_formula_text(self):
        project = example()
        project.table("products").rows[0]["product_name"] = "=not_a_formula"
        project.change(
            project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
        )
        with tempfile.TemporaryDirectory() as folder:
            path = project.export(folder)
            model = parse_export(path)
            products = next(t for t in model.tables if t.name == "products")
            self.assertEqual(products.primary_key.name, "product_id")
            self.assertEqual(products.rows[0]["product_id"], "001")
            self.assertEqual(products.rows[0]["price"], Decimal("350.50"))
            self.assertEqual(products.rows[0]["product_name"], "=not_a_formula")
            self.assertTrue(products.foreign_keys[0].enabled)
            book = load_workbook(
                next(path.glob("prepared_products_*.xlsx")), data_only=False
            )
            self.assertEqual(book.active["B2"].data_type, "s")
            restored = Project.from_dict(
                json.loads((path / "normalizer_project.json").read_text())
            )
            self.assertEqual(restored.as_dict(), project.as_dict())
            self.assertNotEqual(project.export(folder), path)
            with self.assertRaisesRegex(ValueError, "не входят"):
                read_import_profile(path / "import_profile.json", ["wrong.xlsx"])

    def test_large_ids_dates_uuid_project_and_excel_roundtrip(self):
        number = 1234567890123456789
        guid = uuid.uuid4()
        t = table(
            "events",
            [
                ("id", "BIGINT"),
                ("day", "DATE"),
                ("stamp", "TIMESTAMP WITH TIME ZONE"),
                ("guid", "UUID"),
            ],
            [
                dict(
                    id=number,
                    day=date(2026, 2, 1),
                    stamp=datetime(2026, 2, 1, tzinfo=timezone.utc),
                    guid=guid,
                )
            ],
            ("id",),
        )
        project = Project(WorkbookModel([t]))
        restored = Project.from_dict(json.loads(json.dumps(project.as_dict())))
        self.assertEqual(restored.as_dict(), project.as_dict())
        with tempfile.TemporaryDirectory() as folder:
            model = parse_export(project.export(folder))
            row = model.tables[0].rows[0]
            self.assertEqual(row["id"], number)
            self.assertEqual(row["day"], date(2026, 2, 1))
            self.assertEqual(row["stamp"], datetime(2026, 2, 1, tzinfo=timezone.utc))
            self.assertEqual(row["guid"], guid)

    def test_duplicate_and_null_pk_and_missing_parent_block_export(self):
        project = example()
        with self.assertRaises(ValueError):
            project.change(project.set_key, "products", ("manufacturer",))
        project.change(
            project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
        )
        project.table("products").rows[0]["manufacturer_id"] = 999
        with tempfile.TemporaryDirectory() as folder, self.assertRaises(ValueError):
            project.export(folder)
        self.assertEqual(len(project.table("products").rows), 3)

    def test_optional_attributes_and_long_names_survive_export(self):
        project = example()
        project.table("products").rows[0]["country"] = None
        project.table("products").rows[1]["country"] = None
        project.change(
            project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
        )
        self.assertIsNone(project.table("manufacturers").rows[0]["country"])
        long_name = "a" * 60
        project.change(project.rename_table, "manufacturers", long_name)
        t = table("products_data", [("id", "BIGINT")], [dict(id=1)], ("id",))
        project.model.tables.append(t)
        with tempfile.TemporaryDirectory() as folder:
            parsed = parse_export(project.export(folder))
            self.assertEqual(
                set(parsed.table_names()), set(project.model.table_names())
            )
            self.assertTrue(
                next(t for t in parsed.tables if t.name == "products")
                .foreign_keys[0]
                .enabled
            )


@unittest.skipUnless(os.getenv("CHUDO_TEST_DSN"), "Нужен CHUDO_TEST_DSN")
class PreparedPostgresTests(unittest.TestCase):
    def test_import_excel_then_relations_script_is_repeatable_and_enforced(self):
        import psycopg
        from excel2pg.postgres import import_model

        project = example()
        project.change(
            project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
        )
        schema = "prepared_" + uuid.uuid4().hex
        dsn = os.environ["CHUDO_TEST_DSN"]
        try:
            with tempfile.TemporaryDirectory() as folder:
                model = parse_export(project.export(folder, schema))
                result = import_model(model, dsn, schema, add_foreign_keys=False)
                self.assertEqual(result.rows_inserted, 4)
                sql = relations_sql(model, schema)
                with psycopg.connect(dsn, autocommit=True) as c:
                    c.execute(sql)
                    c.execute(sql)
                    self.assertEqual(
                        c.execute(
                            f'SELECT product_id FROM "{schema}".products ORDER BY product_id'
                        ).fetchall(),
                        [("001",), ("002",), ("003",)],
                    )
                    with self.assertRaises(psycopg.errors.ForeignKeyViolation):
                        c.execute(
                            f"UPDATE \"{schema}\".products SET manufacturer_id=999 WHERE product_id='001'"
                        )
                # Профиль тоже создаёт связи. Отдельный SQL после импорта не дублирует их.
                import_model(model, dsn, schema, add_foreign_keys=True)
                with psycopg.connect(dsn, autocommit=True) as c:
                    c.execute(sql)
                    count = c.execute(
                        "SELECT count(*) FROM pg_constraint WHERE contype='f' AND connamespace=%s::regnamespace",
                        (schema,),
                    ).fetchone()[0]
                    self.assertEqual(count, 1)
        finally:
            with psycopg.connect(dsn, autocommit=True) as c:
                c.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')


@unittest.skipUnless(os.getenv("DISPLAY"), "Нужен DISPLAY")
class PreparationUiTests(unittest.TestCase):
    def test_simple_mode_extracts_from_checkboxes_and_auto_ids(self):
        from excelprep.ui import ExcelPreparationApp

        app = ExcelPreparationApp()
        self.addCleanup(app.destroy)
        project = example()
        for c in project.table("products").columns:
            c.is_primary_key = False
        app.loaded(project)
        app.update()
        self.assertEqual(len(app.buttons), 4)
        self.assertEqual(
            app.project.table("products").primary_key.name, "product_row_id"
        )
        self.assertEqual(app.project.table("products").rows[0]["product_id"], "001")
        app.extract_dialog()
        app.update()
        self.assertGreaterEqual(app.extract_canvas.winfo_height(), 150)
        app.extract_fields["manufacturer"].set(True)
        app.extract_create_button.invoke()
        app.update()
        self.assertEqual(app.project.model.table_names(), ["products", "manufacturers"])
        self.assertEqual(
            app.project.table("manufacturers").rows,
            [dict(manufacturer_id=1, manufacturer="Альфа")],
        )
        self.assertEqual(
            [r["manufacturer_id"] for r in app.project.table("products").rows],
            [1, 1, None],
        )
        self.assertEqual(app.project.validate(), [])
        app.toggle_advanced()
        app.update()
        windows = [w for w in app.winfo_children() if w.winfo_class() == "Toplevel"]
        self.assertEqual(len(windows), 1)
        windows[0].destroy()


    def test_real_window_edit_extract_undo_and_links(self):
        from excelprep.ui import ExcelPreparationApp

        app = ExcelPreparationApp()
        self.addCleanup(app.destroy)
        app.loaded(example())
        app.update()
        self.assertEqual(len(app.tree.get_children()), 3)
        ok, _ = app.apply(
            app.project.extract,
            "products",
            ("manufacturer", "country"),
            "manufacturers",
            "manufacturer_id",
        )
        self.assertTrue(ok)
        app.update()
        app.extract_dialog()
        app.update()
        popups = [w for w in app.winfo_children() if w.winfo_class() == "Toplevel"]
        for window in popups:
            window.destroy()
        app.refresh("manufacturers")
        app.update()
        self.assertEqual(len(app.tree.get_children()), 1)
        app.undo()
        app.update()
        self.assertEqual(app.project.model.table_names(), ["products"])
        self.assertEqual(len(app.tree.get_children()), 3)
        app.loaded(example(), append=True)
        app.update()
        self.assertEqual(app.project.model.table_names(), ["products", "products_2"])
        app.undo()
        self.assertEqual(app.project.model.table_names(), ["products"])


if __name__ == "__main__":
    unittest.main()
