import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook


@unittest.skipUnless(os.getenv("DISPLAY"), "Нужен графический экран")
class ProfileUiTests(unittest.TestCase):
    def test_prepared_profile_overrides_old_rules_and_header_setting(self):
        import tkinter as tk
        from interface.import_window import ExcelImportPanel
        root = tk.Tk()
        self.addCleanup(root.destroy)
        panel = ExcelImportPanel(root)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "prepared_products_001.xlsx"
            book = Workbook()
            book.active.append(["product_id", "product_name", "price"])
            book.active.append(["001", "Книга", 100])
            book.save(path)
            profile = {"format": "excelprep-import", "version": 1, "files": [path.name], "table_rules": {path.stem: {"table": "products", "primary_key": ["product_id"], "types": {"product_id": "TEXT", "price": "NUMERIC"}}}, "foreign_keys": []}
            (path.parent / "import_profile.json").write_text(json.dumps(profile))
            with patch("interface.import_window.filedialog.askopenfilenames", return_value=[str(path)]):
                panel.choose_files()
            panel.header_entry.delete(0, "end")
            panel.header_entry.insert(0, "3")
            with patch("config.IMPORT_TABLE_RULES", {"products": {"table": "wrong"}}):
                model = panel._parse()
            self.assertEqual(model.table_names(), ["products"])
            self.assertEqual(model.tables[0].primary_key.name, "product_id")
            self.assertEqual(model.tables[0].rows[0]["product_id"], "001")
