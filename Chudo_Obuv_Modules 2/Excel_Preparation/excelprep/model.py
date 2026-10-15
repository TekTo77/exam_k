"""Переименование, разбиение таблиц и проверка ключей."""

from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
import json
from pathlib import Path
import re
import tempfile
import unicodedata
from uuid import UUID

import pandas as pd

from excel2pg.entities import ColumnEntity, ForeignKeyEntity, TableEntity, WorkbookModel
from excel2pg.types import clean_value

# Знакомые слова переводим; остальные имена транслитерируем.
WORDS = {
    "товары": "products",
    "товар": "product",
    "наименование_товара": "product_name",
    "название_товара": "product_name",
    "пользователи": "users",
    "пользователь": "user",
    "заказы": "orders",
    "заказ": "order",
    "состав_заказа": "order_items",
    "категории": "categories",
    "категория": "category",
    "подкатегории": "subcategories",
    "подкатегория": "subcategory",
    "производители": "manufacturers",
    "производитель": "manufacturer",
    "производство": "manufacturer",
    "бренд": "brand",
    "размеры": "sizes",
    "размер": "size",
    "остатки": "stock_items",
    "склад": "stock_items",
    "роли": "roles",
    "роль": "role",
    "фамилия": "last_name",
    "имя": "first_name",
    "отчество": "middle_name",
    "отчетсво": "middle_name",
    "фио": "client_name",
    "логин": "login",
    "пароль": "password",
    "цена": "price",
    "описание": "description",
    "изображение": "image_file",
    "фото": "image_file",
    "состав": "composition",
    "количество": "quantity",
    "количество_доступное_для_заказа": "quantity_available",
    "дата_заказа": "order_date",
    "номер_заказа": "order_number",
    "цена_за_единицу": "unit_price",
    "название": "name",
    "наименование": "name",
    "код": "id",
    "идентификатор": "id",
    "книги": "books",
    "авторы": "authors",
    "автор": "author",
    "издательства": "publishers",
    "издательство": "publisher",
}
CYRILLIC = dict(
    zip(
        "абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
        (
            "a",
            "b",
            "v",
            "g",
            "d",
            "e",
            "yo",
            "zh",
            "z",
            "i",
            "y",
            "k",
            "l",
            "m",
            "n",
            "o",
            "p",
            "r",
            "s",
            "t",
            "u",
            "f",
            "kh",
            "ts",
            "ch",
            "sh",
            "shch",
            "",
            "y",
            "",
            "e",
            "yu",
            "ya",
        ),
    )
)


def english_name(value, fallback="column"):
    value = re.sub(r"[^\w]+", "_", str(value).strip().lower()).strip("_")
    value = WORDS.get(value, value)
    value = "".join(CYRILLIC.get(c, c) for c in value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^a-z0-9_]+", "_", value).strip("_") or fallback
    if value[0].isdigit():
        value = "_" + value
    return value[:60]


def unique_name(value, used):
    base = english_name(value)
    name, index = base, 2
    while name in used:
        name = f"{base[:52]}_{index}"
        index += 1
    used.add(name)
    return name


def identifier(value):
    value = str(value).strip()
    if not re.fullmatch(r"[a-z_][a-z0-9_]{0,59}", value):
        raise ValueError(
            "Имя: английские буквы, цифры и _, до 60 символов. Начните с буквы или _."
        )
    return value


def singular_name(value):
    return (
        value[:-3] + "y"
        if value.endswith("ies")
        else value[:-1] if value.endswith("s") and not value.endswith("ss") else value
    )


def suggested_names(table, columns):
    """Названия для обычного справочника; их можно поменять в дополнительных настройках."""
    base = singular_name(english_name(columns[0] if columns else "reference"))
    for suffix in ("_name", "_id"):
        if base.endswith(suffix):
            base = base[: -len(suffix)]
    plural = (
        base[:-1] + "ies"
        if base.endswith("y") and not base.endswith(("ay", "ey", "oy", "uy"))
        else base + "s"
    )
    key = base[:45] + "_id"
    occupied = {c.name for c in table.columns}
    if key in occupied:
        key = base[:40] + "_ref_id"
    index = 2
    while key in occupied:
        key = base[:40] + f"_ref_{index}_id"
        index += 1
    return plural[:60], key


def key_values(row, columns):
    return tuple(row.get(c) for c in columns)


def type_for(values):
    values = [v for v in values if v is not None]
    if not values:
        return "TEXT"
    if all(isinstance(v, bool) for v in values):
        return "BOOLEAN"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in values):
        return "BIGINT"
    if all(
        isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) for v in values
    ):
        return "NUMERIC"
    if all(isinstance(v, date) for v in values):
        return "TIMESTAMP" if any(isinstance(v, datetime) for v in values) else "DATE"
    if all(isinstance(v, UUID) for v in values):
        return "UUID"
    return "TEXT"


def load_excel(files, header_row=1):
    from excel2pg.parser import _table_name

    if not isinstance(header_row, int) or header_row < 1:
        raise ValueError("Строка заголовков должна быть числом от 1.")
    model, used = WorkbookModel(), set()
    for path in map(Path, files):
        try:
            with pd.ExcelFile(path, engine="calamine") as book:
                for sheet in book.sheet_names:
                    frame = book.parse(
                        sheet,
                        header=header_row - 1,
                        dtype=object,
                        keep_default_na=False,
                    )
                    headings = list(frame.columns)
                    names, columns, seen = [], [], set()
                    for index, heading in enumerate(headings):
                        if str(heading).startswith("Unnamed:") and all(
                            v == "" for v in frame.iloc[:, index]
                        ):
                            continue
                        name = unique_name(heading, seen)
                        names.append((index, name))
                        columns.append(ColumnEntity(name, str(heading), "TEXT"))
                    rows = [
                        {
                            name: (
                                None
                                if values[index] == ""
                                else clean_value(values[index])
                            )
                            for index, name in names
                        }
                        for values in frame.itertuples(index=False, name=None)
                        if any(
                            values[index] != "" and not pd.isna(values[index])
                            for index, _ in names
                        )
                    ]
                    if not columns:
                        continue
                    for column in columns:
                        column.pg_type = type_for(r[column.name] for r in rows)
                    table = TableEntity(
                        unique_name(_table_name(path.name, str(sheet), set()), used),
                        str(path),
                        str(sheet),
                        columns,
                        rows,
                        header_row=header_row,
                    )
                    singular = (
                        table.name[:-3] + "y"
                        if table.name.endswith("ies")
                        else table.name[:-1] if table.name.endswith("s") else table.name
                    )
                    candidates = [
                        c.name for c in columns if c.name in {"id", f"{singular}_id"}
                    ]
                    for name in candidates:
                        vals = [r[name] for r in rows]
                        if rows and None not in vals and len(set(vals)) == len(rows):
                            table.columns[
                                [c.name for c in columns].index(name)
                            ].is_primary_key = True
                            break
                    model.tables.append(table)
            model.warnings.append(
                f"{path.name}: загружены значения ячеек. Оформление, макросы и формулы не переносятся; формулы должны быть рассчитаны и сохранены в Excel заранее."
            )
        except Exception as error:
            raise ValueError(f"Не удалось прочитать {path.name}: {error}") from error
    if not model.tables:
        raise ValueError("Не найдено листов с заголовками. Проверьте выбранную строку.")
    return model


class Project:
    def __init__(self, model=None):
        self.model = model or WorkbookModel()
        self.history = []

    def table(self, name):
        return next(t for t in self.model.tables if t.name == name)

    def change(self, operation, *args, **kwargs):
        old = deepcopy(self.model)
        try:
            result = operation(*args, **kwargs)
        except Exception:
            self.model = old
            raise
        self.history.append(old)
        self.history = self.history[-10:]
        return result

    def undo(self):
        if self.history:
            self.model = self.history.pop()

    def rename_table(self, old, new):
        new = identifier(new)
        if new != old and new in self.model.table_names():
            raise ValueError("Таблица с таким именем уже есть.")
        self.table(old).name = new
        for table in self.model.tables:
            for fk in table.foreign_keys:
                if fk.referenced_table == old:
                    fk.referenced_table = new

    def rename_column(self, table_name, old, new):
        new = identifier(new)
        table = self.table(table_name)
        if new != old and new in [c.name for c in table.columns]:
            raise ValueError("Колонка с таким именем уже есть.")
        next(c for c in table.columns if c.name == old).name = new
        for row in table.rows:
            row[new] = row.pop(old)
        for child in self.model.tables:
            for fk in child.foreign_keys:
                if child is table:
                    cols = tuple(new if c == old else c for c in fk.columns)
                    fk.column, fk.composite_columns = cols[0], cols
                if fk.referenced_table == table.name:
                    refs = tuple(new if c == old else c for c in fk.reference_columns)
                    fk.referenced_column, fk.composite_references = refs[0], refs

    def set_key(self, table_name, columns):
        table = self.table(table_name)
        self._columns(table, columns)
        if not columns:
            raise ValueError("Выберите хотя бы одну колонку ключа.")
        seen = set()
        for index, row in enumerate(table.rows, 2):
            value = key_values(row, columns)
            if None in value or value in seen:
                raise ValueError(
                    f"Строка {index}: ключ пустой или повторяется: {value}."
                )
            seen.add(value)
        for child in self.model.tables:
            for fk in child.foreign_keys:
                if fk.referenced_table == table_name and fk.reference_columns != tuple(
                    columns
                ):
                    raise ValueError(
                        "На старый ключ уже есть связь. Сначала удалите её в окне связей."
                    )
        for c in table.columns:
            c.is_primary_key = c.name in columns

    @staticmethod
    def _columns(table, columns):
        if len(set(columns)) != len(columns) or not set(columns) <= {
            c.name for c in table.columns
        }:
            raise ValueError("Выберите существующие колонки без повторов.")

    def add_id(self, table_name, name):
        table, name = self.table(table_name), identifier(name)
        if name in {c.name for c in table.columns}:
            raise ValueError(
                "Колонка уже существует; выберите её как ключ или задайте другое имя."
            )
        table.columns.insert(0, ColumnEntity(name, "(created)", "BIGINT", False))
        for index, row in enumerate(table.rows, 1):
            row[name] = index
        self.set_key(table_name, (name,))

    def prepare_keys(self):
        created = []
        for table in self.model.tables:
            if table.primary_keys:
                continue
            used = {c.name for c in table.columns}
            name = singular_name(table.name)[:45] + "_id"
            if name in used:
                name = singular_name(table.name)[:40] + "_row_id"
            name = unique_name(name, used)
            self.add_id(table.name, name)
            created.append((table.name, name))
        return created

    def extract(self, source, columns, target, id_column, match_columns=()):
        table = self.table(source)
        target, id_column = identifier(target), identifier(id_column)
        self._columns(table, columns)
        explicit_match = bool(match_columns)
        match_columns = tuple(match_columns or columns)
        self._columns(table, match_columns)
        if not columns or not set(match_columns) <= set(columns):
            raise ValueError("Ключ совпадения должен состоять из выносимых колонок.")
        if target in self.model.table_names():
            raise ValueError("Имя новой таблицы уже занято.")
        if id_column in {c.name for c in table.columns}:
            raise ValueError("Имя нового ID уже занято в исходной таблице.")
        if any(c.is_primary_key and c.name in columns for c in table.columns):
            raise ValueError(
                "Не выносите первичный ключ строки. Для разбиения заказа сначала создайте отдельный ID позиции."
            )
        moved = []
        for fk in table.foreign_keys:
            overlap = set(fk.columns) & set(columns)
            if overlap and overlap != set(fk.columns):
                raise ValueError("Составной внешний ключ нужно выносить целиком.")
            if overlap:
                moved.append(fk)
        for other in self.model.tables:
            if any(
                fk.referenced_table == source
                and set(fk.reference_columns) & set(columns)
                for fk in other.foreign_keys
            ):
                raise ValueError("На выносимые колонки уже ссылается другая таблица.")
        records, identities, row_ids = [], {}, []
        for row in table.rows:
            data = key_values(row, columns)
            if all(v is None for v in data):
                row_ids.append(None)
                continue
            key = key_values(row, match_columns)
            if explicit_match and any(v is None for v in key):
                raise ValueError(
                    f"Пустой ключ совпадения {key}. Выберите другой набор колонок."
                )
            if key in identities:
                index, previous = identities[key]
                if previous != data:
                    raise ValueError(
                        f"Для {key} разные значения {data} и {previous}. Исправьте данные или расширьте ключ совпадения."
                    )
            else:
                index = len(records) + 1
                identities[key] = index, data
                records.append({id_column: index, **dict(zip(columns, data))})
            row_ids.append(index)
        parent_columns = [
            ColumnEntity(id_column, "(created)", "BIGINT", False, True)
        ] + [deepcopy(c) for c in table.columns if c.name in columns]
        parent = TableEntity(
            target, "(prepared)", "Sheet1", parent_columns, records, deepcopy(moved)
        )
        table.columns = [c for c in table.columns if c.name not in columns]
        table.columns.append(ColumnEntity(id_column, "(created)", "BIGINT"))
        for row, index in zip(table.rows, row_ids):
            for col in columns:
                row.pop(col, None)
            row[id_column] = index
        table.foreign_keys = [fk for fk in table.foreign_keys if fk not in moved]
        self.model.tables.append(parent)
        self.add_relation(source, (id_column,), target, (id_column,))
        return target

    def add_relation(self, child_name, columns, parent_name, references):
        child, parent = self.table(child_name), self.table(parent_name)
        self._columns(child, columns)
        self._columns(parent, references)
        if not columns or len(columns) != len(references):
            raise ValueError("Количество колонок с обеих сторон должно совпадать.")
        if tuple(references) != tuple(c.name for c in parent.primary_keys):
            raise ValueError(
                "Родительские колонки должны быть первичным ключом в том же порядке."
            )
        present = {key_values(r, references) for r in parent.rows}
        if len(present) != len(parent.rows) or any(None in key for key in present):
            raise ValueError(
                "Ключ родительской таблицы пустой или повторяется. Исправьте его сначала."
            )
        orphans = {
            key_values(r, columns)
            for r in child.rows
            if all(r.get(c) is not None for c in columns)
        } - present
        if orphans:
            raise ValueError(f"Не найдены родительские значения: {list(orphans)[:5]}.")
        child.foreign_keys = [
            f for f in child.foreign_keys if f.columns != tuple(columns)
        ]
        child.foreign_keys.append(
            ForeignKeyEntity(
                columns[0],
                parent_name,
                references[0],
                "explicit",
                0,
                True,
                tuple(columns),
                tuple(references),
            )
        )

    def link_by_values(
        self, child_name, columns, parent_name, match_columns, id_column
    ):
        child, parent = self.table(child_name), self.table(parent_name)
        self._columns(child, columns)
        self._columns(parent, match_columns)
        if not columns or len(columns) != len(match_columns):
            raise ValueError("Выберите одинаковое количество колонок для совпадения.")
        if len(parent.primary_keys) != 1:
            raise ValueError(
                "Замена названия на ID требует одного первичного ключа родителя."
            )
        id_column = identifier(id_column)
        if id_column in {c.name for c in child.columns}:
            raise ValueError("Имя нового столбца ID уже занято.")
        lookup = {}
        for row in parent.rows:
            key = key_values(row, match_columns)
            if key in lookup:
                raise ValueError(
                    f"Неоднозначное совпадение {key}. Добавьте колонку для различения."
                )
            lookup[key] = row[parent.primary_key.name]
        values = []
        for row in child.rows:
            key = key_values(row, columns)
            if all(v is None for v in key):
                values.append(None)
            elif key not in lookup:
                raise ValueError(
                    f"Не найдено совпадение {key}. Исправьте исходные данные."
                )
            else:
                values.append(lookup[key])
        child.columns.append(
            ColumnEntity(id_column, "(created)", parent.primary_key.pg_type)
        )
        for row, value in zip(child.rows, values):
            row[id_column] = value
        self.add_relation(
            child_name, (id_column,), parent_name, (parent.primary_key.name,)
        )

    def preview_link(
        self,
        child_name,
        columns,
        parent_name,
        match_columns,
        reference_column,
        mode="text",
        combine=False,
        normalize=True,
        replace_existing=False,
    ):
        """Сначала показываем совпадения. ФИО можно хранить одним полем или частями."""
        import unicodedata

        child, parent = self.table(child_name), self.table(parent_name)
        self._columns(child, columns)
        self._columns(parent, match_columns)
        if not columns or not match_columns:
            raise ValueError("Выберите поля, по которым таблицы связаны.")
        if len(parent.primary_keys) != 1:
            raise ValueError(
                "У родителя составной ключ. Используйте Ещё → Составная связь."
            )
        reference_column = identifier(reference_column)
        if reference_column in {c.name for c in child.primary_keys}:
            raise ValueError(
                "Ссылка не должна заменять ID строки. Задайте другое имя для новой ссылки."
            )
        if mode not in {"id", "text"}:
            raise ValueError("Выберите сравнение по ID или по тексту.")
        if not combine and len(columns) != len(match_columns):
            raise ValueError(
                "Число полей должно совпадать. Для ФИО включите объединение полей."
            )
        if mode == "id" and combine:
            raise ValueError("Для ID объединение текстовых полей не используется.")

        def value_key(value, pg_type=None):
            if value is None or value == "":
                return None
            if mode == "id":
                if pg_type in {
                    "SMALLINT",
                    "INTEGER",
                    "BIGINT",
                    "NUMERIC",
                    "DOUBLE PRECISION",
                    "REAL",
                } or (pg_type or "").startswith("NUMERIC("):
                    try:
                        number = Decimal(str(value).strip())
                        return number if number.is_finite() else ("invalid", str(value))
                    except ArithmeticError:
                        return ("invalid", str(value))
                if pg_type == "UUID":
                    try:
                        return UUID(str(value).strip())
                    except ValueError:
                        return ("invalid", str(value))
                return str(value)
            text = unicodedata.normalize("NFKC", str(value))
            return " ".join(text.split()).casefold() if normalize else text

        types = {c.name: c.pg_type for c in parent.columns}

        def row_key(row, fields):
            values = [row.get(c) for c in fields]
            if all(v is None or v == "" for v in values):
                return None
            if combine:
                return value_key(
                    " ".join(str(v) for v in values if v is not None and v != "")
                )
            return tuple(
                value_key(v, types.get(match_columns[i])) for i, v in enumerate(values)
            )

        parent_keys = [r.get(parent.primary_key.name) for r in parent.rows]
        if None in parent_keys or len(set(parent_keys)) != len(parent_keys):
            raise ValueError(
                "ID родительской таблицы пустой или повторяется. Исправьте ключ сначала."
            )
        lookup = {}
        for row in parent.rows:
            key = row_key(row, match_columns)
            if key is not None:
                lookup.setdefault(key, []).append(row[parent.primary_key.name])

        # Готовые правильные ID сохраняем, даже если текст отсутствует.
        def pk_token(value):
            if value is None or value == "":
                return None
            if parent.primary_key.pg_type in {
                "BIGINT",
                "INTEGER",
                "SMALLINT",
                "NUMERIC",
                "DOUBLE PRECISION",
                "REAL",
            } or parent.primary_key.pg_type.startswith("NUMERIC("):
                try:
                    number = Decimal(str(value))
                    return number if number.is_finite() else str(value)
                except ArithmeticError:
                    return str(value)
            if parent.primary_key.pg_type == "UUID":
                try:
                    return UUID(str(value))
                except ValueError:
                    return str(value)
            return str(value)

        known_ids = {pk_token(k): k for k in parent_keys}
        report = {
            "matched": 0,
            "kept": 0,
            "empty": 0,
            "missing": 0,
            "ambiguous": 0,
            "conflicts": 0,
            "issues": [],
            "values": [],
        }
        for index, row in enumerate(child.rows, 2):
            key = row_key(row, columns)
            matches = lookup.get(key, []) if key is not None else []
            old = row.get(reference_column)
            old_id = known_ids.get(pk_token(old)) if old is not None else None
            issue = None
            if len(matches) == 1:
                value = matches[0]
                report["matched"] += 1
                if (
                    old is not None
                    and old != ""
                    and pk_token(old) != pk_token(value)
                    and not replace_existing
                ):
                    report["conflicts"] += 1
                    issue = f"Строка {index}: уже записан ID {old}, по выбранным полям получается {value}."
            elif old_id is not None and not replace_existing:
                value = old_id
                report["kept"] += 1
            else:
                value = None
                if key is None:
                    report["empty"] += 1
                elif matches:
                    report["ambiguous"] += 1
                    issue = f"Строка {index}: {key!r} соответствует нескольким записям {parent.name}: {matches[:4]}."
                else:
                    report["missing"] += 1
                    issue = f"Строка {index}: {key!r} не найдено в {parent.name}."
            report["values"].append(value)
            if issue and len(report["issues"]) < 8:
                report["issues"].append(issue)
        return report

    def connect_tables(
        self,
        child_name,
        columns,
        parent_name,
        match_columns,
        reference_column,
        mode="text",
        combine=False,
        normalize=True,
        replace_existing=False,
        allow_missing=False,
    ):
        report = self.preview_link(
            child_name,
            columns,
            parent_name,
            match_columns,
            reference_column,
            mode,
            combine,
            normalize,
            replace_existing,
        )
        if report["conflicts"] or (
            (report["missing"] or report["ambiguous"]) and not allow_missing
        ):
            raise ValueError("Связь пока не создана.\n" + "\n".join(report["issues"]))
        child, parent = self.table(child_name), self.table(parent_name)
        column = next((c for c in child.columns if c.name == reference_column), None)
        if column is None:
            column = ColumnEntity(
                reference_column, "(resolved reference)", parent.primary_key.pg_type
            )
            child.columns.append(column)
        if any(v is None for v in report["values"]) and not column.nullable:
            raise ValueError(
                "В этой колонке запрещены пустые ссылки. Исправьте отсутствующие совпадения."
            )
        column.pg_type = parent.primary_key.pg_type
        for row, value in zip(child.rows, report["values"]):
            row[reference_column] = value
        self.add_relation(
            child_name, (reference_column,), parent_name, (parent.primary_key.name,)
        )
        errors = self.validate()
        if errors:
            raise ValueError(
                "Связь нарушает существующую структуру:\n" + "\n".join(errors[:5])
            )
        return report

    def validate(self):
        errors = []
        if len(set(self.model.table_names())) != len(self.model.tables):
            errors.append("Имена таблиц повторяются.")
        for table in self.model.tables:
            if not table.columns or len({c.name for c in table.columns}) != len(
                table.columns
            ):
                errors.append(f"{table.name}: нет колонок или их имена повторяются.")
            keys = tuple(c.name for c in table.primary_keys)
            if not keys:
                errors.append(f"{table.name}: не задан первичный ключ.")
            else:
                seen = set()
                for row in table.rows:
                    value = key_values(row, keys)
                    if None in value or value in seen:
                        errors.append(
                            f"{table.name}: пустой или повторяющийся ключ {value}."
                        )
                        break
                    seen.add(value)
            for fk in table.foreign_keys:
                try:
                    parent = self.table(fk.referenced_table)
                    if fk.reference_columns != tuple(
                        c.name for c in parent.primary_keys
                    ):
                        raise ValueError("родительский ключ изменился")
                    present = {key_values(r, fk.reference_columns) for r in parent.rows}
                    missing = {
                        key_values(r, fk.columns)
                        for r in table.rows
                        if all(r.get(c) is not None for c in fk.columns)
                    } - present
                    if missing:
                        raise ValueError(
                            f"нет родительских строк для {list(missing)[:3]}"
                        )
                    child_types = [
                        next(c.pg_type for c in table.columns if c.name == n)
                        for n in fk.columns
                    ]
                    parent_types = [
                        next(c.pg_type for c in parent.columns if c.name == n)
                        for n in fk.reference_columns
                    ]
                    for left, right in zip(child_types, parent_types):
                        if left != right and not {left, right} <= {
                            "BIGINT",
                            "INTEGER",
                            "SMALLINT",
                            "NUMERIC",
                            "DOUBLE PRECISION",
                        }:
                            raise ValueError(f"разные типы ключей: {left}, {right}")
                except (ValueError, StopIteration) as error:
                    errors.append(f"{table.name} → {fk.referenced_table}: {error}.")
        return errors

    def suggestions(self):
        lines = []
        for table in self.model.tables:
            for col in table.columns:
                if col.is_primary_key or any(
                    col.name in f.columns for f in table.foreign_keys
                ):
                    continue
                vals = [
                    r.get(col.name) for r in table.rows if r.get(col.name) is not None
                ]
                if vals and len(set(vals)) < len(vals) and col.pg_type == "TEXT":
                    lines.append(
                        f"{table.name}.{col.name}: {len(set(vals))} разных значений в {len(vals)} строках. Возможно, нужен справочник; проверьте смысл данных."
                    )
        return lines

    def save_project(self, path):
        Path(path).write_text(
            json.dumps(self.as_dict(), ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def as_dict(self):
        from dataclasses import asdict

        return {
            "format": "excelprep",
            "version": 1,
            "model": encode(asdict(self.model)),
        }

    @classmethod
    def from_dict(cls, data):
        if data.get("format") != "excelprep" or data.get("version") != 1:
            raise ValueError(
                "Это не проект редактора Excel или версия не поддерживается."
            )
        raw = decode(data["model"])
        tables = []
        for t in raw["tables"]:
            t["columns"] = [ColumnEntity(**c) for c in t["columns"]]
            t["foreign_keys"] = [
                ForeignKeyEntity(
                    **{
                        **f,
                        "composite_columns": tuple(f["composite_columns"]),
                        "composite_references": tuple(f["composite_references"]),
                    }
                )
                for f in t["foreign_keys"]
            ]
            identifier(t["name"])
            for c in t["columns"]:
                identifier(c.name)
            tables.append(TableEntity(**t))
        return cls(WorkbookModel(tables, raw.get("warnings", [])))

    def export(self, directory, schema="public"):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        errors = self.validate()
        if errors or not self.model.tables:
            raise ValueError(
                "Сначала исправьте структуру:\n" + "\n".join(errors or ["Нет таблиц."])
            )
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        rules, relations, files = {}, [], []
        # Новая папка при каждом сохранении: исходные файлы не перезаписываем.
        stage = Path(tempfile.mkdtemp(prefix=".preparing_", dir=directory))
        try:
            for index, table in enumerate(self.model.tables, 1):
                # Короткое имя не обрежется в импортёре; номер исключает коллизии.
                file_stem = f"prepared_{table.name[:42]}_{index:03}"
                file_name = file_stem + ".xlsx"
                files.append(file_name)
                book = Workbook()
                sheet = book.active
                sheet.title = "Sheet1"
                sheet.append([c.name for c in table.columns])
                for row in table.rows:
                    values = []
                    for col in table.columns:
                        value = row.get(col.name)
                        if isinstance(value, (Decimal, UUID)) or (
                            isinstance(value, int) and abs(value) >= 10**15
                        ):
                            value = str(value)
                        if isinstance(value, datetime) and value.tzinfo:
                            value = value.isoformat()
                        values.append(value)
                    sheet.append(values)
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="285A71")
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for index, col in enumerate(table.columns, 1):
                    sheet.column_dimensions[get_column_letter(index)].width = min(
                        48, max(16, len(col.name) + 3)
                    )
                book.save(stage / file_name)
                rules[file_stem] = {
                    "table": table.name,
                    "primary_key": [c.name for c in table.primary_keys],
                    "types": {c.name: c.pg_type for c in table.columns},
                }
                for fk in table.foreign_keys:
                    relations.append(
                        {
                            "table": table.name,
                            "columns": list(fk.columns),
                            "references": fk.referenced_table,
                            "reference_columns": list(fk.reference_columns),
                        }
                    )
            profile = {
                "format": "excelprep-import",
                "version": 1,
                "files": files,
                "table_rules": rules,
                "foreign_keys": relations,
                "auto_foreign_keys": False,
                "reference_rules": [],
            }
            (stage / "import_profile.json").write_text(
                json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            config = (
                "# Необязательно: эти настройки можно перенести в config.py.\n"
                + "IMPORT_TABLE_RULES = "
                + repr(rules)
                + "\nIMPORT_FOREIGN_KEYS = "
                + repr(tuple(relations))
                + "\nIMPORT_AUTO_FOREIGN_KEYS = False\nIMPORT_REFERENCE_RULES = ()\n"
            )
            (stage / "import_settings.txt").write_text(config, encoding="utf-8")
            (stage / "table_relations.sql").write_text(
                relations_sql(self.model, schema), encoding="utf-8"
            )
            self.save_project(stage / "normalizer_project.json")
            name = "prepared_" + datetime.now().strftime("%Y%m%d_%H%M%S")
            target, count = directory / name, 2
            while target.exists():
                target = directory / f"{name}_{count}"
                count += 1
            stage.rename(target)
            return target
        except Exception:
            import shutil

            shutil.rmtree(stage)
            raise


def encode(value):
    if isinstance(value, UUID):
        return {"__type__": "uuid", "value": str(value)}
    if isinstance(value, datetime):
        return {"__type__": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"__type__": "date", "value": value.isoformat()}
    if isinstance(value, Decimal):
        return {"__type__": "decimal", "value": str(value)}
    if isinstance(value, dict):
        return {k: encode(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [encode(v) for v in value]
    return value


def decode(value):
    if isinstance(value, dict):
        if value.get("__type__") in {"date", "datetime", "decimal", "uuid"} and set(
            value
        ) == {"__type__", "value"}:
            return {
                "date": date.fromisoformat,
                "datetime": datetime.fromisoformat,
                "decimal": Decimal,
                "uuid": UUID,
            }[value["__type__"]](value["value"])
        return {k: decode(v) for k, v in value.items()}
    if isinstance(value, list):
        return [decode(v) for v in value]
    return value


def relations_sql(model, schema="public"):
    from excel2pg.constraint_names import _constraint_name

    quote = lambda v: '"' + str(v).replace('"', '""') + '"'
    statements = [
        "-- Только отношения. Сначала импортируйте подготовленные Excel с import_profile.json.",
        "BEGIN;",
    ]
    for table in model.tables:
        for fk in table.foreign_keys:
            name = _constraint_name(table, fk)
            child = f"{quote(schema)}.{quote(table.name)}"
            parent = f"{quote(schema)}.{quote(fk.referenced_table)}"
            cols = ", ".join(map(quote, fk.columns))
            refs = ", ".join(map(quote, fk.reference_columns))
            child_literal = child.replace("'", "''")
            tag = "$relation$"
            while tag in child + parent:
                tag = tag[:-1] + "x$"
            statements.append(
                f"DO {tag} BEGIN\n  IF NOT EXISTS (SELECT 1 FROM pg_constraint WHERE conname = '{name}' AND conrelid = '{child_literal}'::regclass) THEN\n    ALTER TABLE {child} ADD CONSTRAINT {quote(name)} FOREIGN KEY ({cols}) REFERENCES {parent} ({refs});\n  END IF;\nEND {tag};"
            )
    return "\n\n".join(statements + ["COMMIT;"]) + "\n"
